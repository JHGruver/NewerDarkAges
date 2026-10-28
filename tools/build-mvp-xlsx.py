#!/usr/bin/env python3
"""Builds the two V2 workbooks from MVP-V2.json (written by tools/build-mvp-v2.cjs):
  NDA_MVP_Characters_V2.xlsx  the MVP roster and every meta detail, prompts, tracker, open slots, vote scores
  NDA_Art_Color_Reference.xlsx  a reference for color balancing and inspiration (palette rules, checks on every MVP trio, skin chart, inspiration)
Needs the project venv (tools/.venv with openpyxl)."""
import json, os, colorsys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

root = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
D = json.load(open(os.path.join(root, "MVP-V2.json")))
SL = D["slots"]; REAL = [s for s in SL if s.get("id")]; A = D["art2"]; PEO = {p["id"]: p for p in D["peoples"]}
HDR = PatternFill("solid", fgColor="1F1F2B"); HF = Font(bold=True, color="FFFFFF"); thin = Side(style="thin", color="DDDDDD")
def fill(h): return PatternFill("solid", fgColor=h.lstrip("#").upper())
def rgb(h): h = h.lstrip("#"); return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))
def lin(c): c = c / 255; return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
def lum(h): r, g, b = (lin(v) for v in rgb(h)); return 0.2126 * r + 0.7152 * g + 0.0722 * b
def contrast(a, b): la, lb = sorted((lum(a), lum(b)), reverse=True); return (la + 0.05) / (lb + 0.05)
def hls(h): r, g, b = (v / 255 for v in rgb(h)); hh, l, s = colorsys.rgb_to_hls(r, g, b); return hh * 360, s, l
def hdist(a, b): d = abs(a - b) % 360; return min(d, 360 - d)
def txt(h): return "000000" if lum(h) > 0.3 else "FFFFFF"
def style(ws, widths, hdr_row=1, wrap=(), swatch=(), freeze="C2"):
    for c in ws[hdr_row]: c.fill = HDR; c.font = HF; c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=hdr_row + 1):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=((c.column - 1) in wrap)); c.border = Border(bottom=thin)
            if (c.column - 1) in swatch and isinstance(c.value, str) and c.value.startswith("#"):
                c.fill = fill(c.value); c.font = Font(color=txt(c.value), bold=True)
    if freeze: ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
def banner(ws, rows, widths):
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value and str(ws.cell(r, 1).value).isupper(): ws.cell(r, 1).font = Font(bold=True, size=13)
        if ws.cell(r, 1).value and ws.cell(r, 2).value is not None and not ws.cell(r, 1).value.isupper(): ws.cell(r, 1).font = Font(bold=True)

# =============================================================================================== MVP CHARACTERS V2
wb = Workbook(); ws = wb.active; ws.title = "START HERE"
nart = sum(1 for s in REAL if s["v1art"])
banner(ws, [
    ["NEWER DARK AGES: MVP CHARACTERS, VERSION 2", "The launch roster only: 24 slots, 8 peoples, 3 each. Generated from the site data, so it never drifts. Rebuild: node tools/build-mvp-v2.cjs"],
    [],
    ["What is in it", ""],
    ["MVP Roster", "One row per slot with every meta detail: class, type, body, face, costume, signature exaggeration, abilities, lore, colors, notes."],
    ["Prompts", "The art direction V2 (Shard Cel) character and background prompts, ready to paste. Character count and the style check for each."],
    ["Production Tracker", "Fill in as art is made. Dropdowns for status, plus the exact file name to save each image as (art/shard/). The site finds the files on its own."],
    ["Open Slots", "The slots still to decide, with candidates and the vote scores for each."],
    ["Vote Scores", "What Jacob and Mike scored every MVP character in the vote."],
    ["Changes", "What changed from the full roster to get here."],
    [],
    ["Status", f"{len(REAL)} chosen characters, {len(SL) - len(REAL)} open slots. {nart} chosen characters have art in the ORIGINAL style (to be redrawn). 0 have V2 art yet."],
    ["Art direction", "Shard Cel. See ART-DIRECTION-V2.md and the Art Direction page. No prompt names an artist or a show."],
    ["Colors", "Each character keeps its three accent colors (main 60%, secondary 30%, pop 10%). The people color is the background wash only. See NDA_Art_Color_Reference.xlsx for the balance checks."],
], [24, 120])

ws = wb.create_sheet("MVP Roster")
H = ["#", "People", "Slot", "Character", "ID", "Class", "Type", "Role", "Size", "Movement", "Physique", "Age", "Gender", "Heritage / source", "Identity", "Face", "Outfit and gear (art spec)", "Signature exaggeration (V2)", "Pose", "Tagline", "Summary", "Strengths", "Weaknesses", "Abilities", "Lore", "Main", "Secondary", "Pop", "People color", "New or changed", "Notes", "Original-style art?", "V2 art"]
ws.append(H)
for s in SL:
    if not s.get("id"):
        ws.append([s["n"], PEO[s["people"]]["name"], s["slot"], "TBD: " + s["tbd"], "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", s["ask"]] + [""] * 11 + ["TBD"]); continue
    ws.append([s["n"], PEO[s["people"]]["name"], s["slot"], s["name"], s["id"], s["cls"], s["type"], s["role"], s["size"], s["move"], s["physique"], s["age"], s["gender"], s["ethnicity"], s["who"], s["expression"], s["outfit"], s["big"], s["pose"], s["tagline"], s["summary"], "\n".join("• " + x for x in s["strengths"]), "\n".join("• " + x for x in s["weaknesses"]), "\n".join("• " + x for x in s["abilities"]), s["lore"], s["pal"][0], s["pal"][1], s["pal"][2], s["peopleColor"], "New" if s["isNew"] else ("Redesigned" if s["redesign"] else ""), s["note"], "Yes, redraw" if s["v1art"] else "No", "Not yet"])
style(ws, [4, 14, 5, 28, 18, 18, 11, 9, 8, 11, 30, 9, 9, 26, 34, 28, 52, 46, 44, 30, 46, 36, 36, 46, 52, 10, 10, 10, 11, 12, 40, 14, 10], wrap=(10, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 30), swatch=(25, 26, 27, 28), freeze="E2")

ws = wb.create_sheet("Prompts")
ws.append(["#", "People", "Character", "Style check", "Char prompt length", "Character prompt (3:4, transparent PNG)", "Background prompt (16:9)", "Main", "Secondary", "Pop", "People color"])
for s in REAL:
    ws.append([s["n"], PEO[s["people"]]["name"], s["name"], "Passed" if not s["lint"] else "; ".join(s["lint"]), len(s["charPrompt"]), s["charPrompt"], s["bgPrompt"], s["pal"][0], s["pal"][1], s["pal"][2], s["peopleColor"]])
style(ws, [4, 14, 28, 14, 10, 110, 90, 10, 10, 10, 11], wrap=(5, 6), swatch=(7, 8, 9, 10), freeze="D2")
for r in range(2, ws.max_row + 1): ws.row_dimensions[r].height = 160
ws.conditional_formatting.add(f"D2:D{ws.max_row}", CellIsRule(operator="equal", formula=['"Passed"'], fill=fill("C6EFCE")))
ws.conditional_formatting.add(f"D2:D{ws.max_row}", CellIsRule(operator="notEqual", formula=['"Passed"'], fill=fill("FFC7CE")))

ws = wb.create_sheet("Production Tracker")
ws.append(["#", "People", "Character", "Character art (V2)", "Background art (V2)", "Matches Shard Cel test?", "Reviewer", "Date", "Notes", "Original-style art exists?", "Save portrait as", "Save background as"])
for s in REAL: ws.append([s["n"], PEO[s["people"]]["name"], s["name"], "Not started", "Not started", "", "", "", "", "Yes" if s["v1art"] else "No", f"art/shard/{s['id']}.png", f"art/shard/{s['id']}-bg.jpg"])
style(ws, [4, 14, 30, 18, 18, 22, 14, 12, 50, 18, 34, 38], wrap=(8,), freeze="D2")
st = DataValidation(type="list", formula1='"Not started,Prompt ready,Generating,Needs redo,Approved,Wired into site"', allow_blank=True); ws.add_data_validation(st); st.add(f"D2:E{ws.max_row}")
ok = DataValidation(type="list", formula1='"Yes,No,Partly"', allow_blank=True); ws.add_data_validation(ok); ok.add(f"F2:F{ws.max_row}")
for col in ("D", "E"):
    ws.conditional_formatting.add(f"{col}2:{col}{ws.max_row}", CellIsRule(operator="equal", formula=['"Approved"'], fill=fill("C6EFCE")))
    ws.conditional_formatting.add(f"{col}2:{col}{ws.max_row}", CellIsRule(operator="equal", formula=['"Wired into site"'], fill=fill("9BD7A8")))
    ws.conditional_formatting.add(f"{col}2:{col}{ws.max_row}", CellIsRule(operator="equal", formula=['"Needs redo"'], fill=fill("FFC7CE")))
    ws.conditional_formatting.add(f"{col}2:{col}{ws.max_row}", CellIsRule(operator="equal", formula=['"Generating"'], fill=fill("FFEB9C")))

ws = wb.create_sheet("Open Slots")
ws.append(["People", "Slot", "What is open", "Question", "Candidate", "Jacob", "Mike", "Has original-style art?"])
names = D["voters"]; lab = {"-1": "Dislike", "1": "Like", "2": "Super like"}
for s in SL:
    if s.get("id"): continue
    for c in s["candidates"] or [{"name": "(a new concept)", "v": [], "art": False}]:
        ws.append([PEO[s["people"]]["name"], s["slot"], s["tbd"], s["ask"], c["name"], lab.get(str(c["v"][0]), "") if len(c["v"]) > 0 else "", lab.get(str(c["v"][1]), "") if len(c["v"]) > 1 else "", ""])
style(ws, [14, 6, 30, 56, 30, 12, 12, 20], wrap=(2, 3), freeze="A2")

ws = wb.create_sheet("Vote Scores")
ws.append(["#", "People", "Character"] + names + ["Combined (average)"])
for s in REAL:
    v = s["votes"]; ws.append([s["n"], PEO[s["people"]]["name"], s["name"]] + [lab.get(str(x), "") for x in v] + [round(sum(x for x in v if x is not None) / max(1, len([x for x in v if x is not None])), 2)])
style(ws, [4, 14, 32, 12, 12, 18], freeze="D2")

ws = wb.create_sheet("Changes")
ws.append(["Change", "Detail"])
for c in D["changes"]: ws.append([c["t"], c["d"]])
ws.append(["", ""]); ws.append(["STILL OPEN", ""])
for o in D["open"]: ws.append([o, ""])
style(ws, [48, 110], wrap=(1,), freeze="A2")
wb.save(os.path.join(root, "NDA_MVP_Characters_V2.xlsx"))

# =============================================================================================== COLOR REFERENCE
wb = Workbook(); ws = wb.active; ws.title = "START HERE"
banner(ws, [
    ["NEWER DARK AGES: ART AND COLOR REFERENCE (V2)", "A reference for color balancing and inspiration. It goes with the Shard Cel art direction and the MVP character workbook."],
    [],
    ["Color Rules", "The balance rules and the thresholds the checks use."],
    ["People Palettes", "The eight background wash colors, with mood, harmony suggestions and what to avoid."],
    ["MVP Trios", "Every MVP character's three accent colors, checked against the rules, with flags and ideas. Fix the flags first."],
    ["Skin Chart", "Flat skin tones for humanoids so skin never competes with the three accents."],
    ["Shape and Material", "The shapes and drawing shortcuts that keep the style consistent."],
    ["Inspiration", "Per people: mood, shapes, light and scene ideas. Per element: what to look for when gathering references (described by qualities, not by names)."],
    [],
    ["How to use", "Before drawing a character, read its row in MVP Trios. If it has a flag, adjust the color in the prompt, then rerun tools/build-mvp-v2.cjs."],
], [22, 120])

ws = wb.create_sheet("Color Rules")
ws.append(["Rule", "What it means", "Check or threshold"])
for r in [
    ["60 / 30 / 10", "Main about 60% of the colored area, secondary about 30%, pop about 10%.", "Every character has exactly three accents."],
    ["Black, white and grey are free", "Outlines, spot blacks, bone, metal and cloth highlights are not counted as accents.", f"Outline color {A['outline']}; paper white {A['white']}."],
    ["People color is the background", "The people color washes the scene (about 65%). It is never used on the character.", "Hue distance between the people color and each accent: at least 25 degrees, or the figure melts into its scene."],
    ["Value spread", "A character must read in grayscale. One dark, one light, one in between.", "Brightest accent minus darkest accent: relative luminance difference of at least 0.25."],
    ["Main versus secondary", "The two big colors need to separate cleanly.", "Contrast ratio of at least 1.5, or a hue gap of at least 40 degrees."],
    ["Pop must pop", "The 10% color is the loudest: the most saturated or the strongest contrast to main.", "Pop saturation at least 0.55, or contrast to main of at least 2."],
    ["No muddy mid-tones", "Avoid accents that are both low-saturation and mid-lightness.", "Flag when saturation is below 0.18 and lightness is between 0.25 and 0.75 (unless it is meant as a neutral)."],
    ["Skin from the chart", "Humanoid skin is one flat tone with one shadow tone, from the Skin Chart.", "Do not tint skin with an accent."],
    ["Shadow tone", "Each color has one darker hard-edged shadow tone: same hue, about 30% darker, slightly more saturated.", "No grey shadows."],
    ["Two tones per material", "Base plus shadow. Pure black only for spot fills in the deepest creases.", "No third tone, no gradients."],
]: ws.append(r)
style(ws, [26, 90, 90], wrap=(1, 2), freeze="A2")

ws = wb.create_sheet("People Palettes")
ws.append(["People", "Wash color", "Name", "Hue", "Saturation", "Lightness", "Complement", "Split complement A", "Split complement B", "Analogous A", "Analogous B", "Mood", "Light story", "Avoid on the figure"])
MOOD = {
    "human": ("scrappy, rain-slick, sodium-lamp warmth, hand-built", "warm gold street lamps against wet asphalt", "gold or yellow-orange that matches the wash"),
    "vampire": ("opulent decay, rooftop night, gothic glamour", "pale violet rim light on near-black", "pure black bodies with no accent to break them up"),
    "were": ("scrapyard loyalty, feral, tribal, burnt", "orange campfire light on rust", "rust orange or brown that matches the wash"),
    "ghost": ("static, cold fog, faded and flickering", "cyan CRT glow in blue-white fog", "ice blue that vanishes into the wash"),
    "undead": ("bureaucratic mortality, dust, bone and stamps", "warm dusty light on pale grey-beige", "beige and bone as the main color"),
    "forest": ("overgrown suburbia, mossy, patient", "green-tinted light under leaves", "moss green as the main color"),
    "eldritch": ("wrong geometry, flooded cathedral, vast", "dim teal from below, one sickly highlight", "teal or deep green as the main color"),
    "magic": ("ritual precinct, learned and strange", "violet glow from chalk and spore light", "violet as the main color"),
}
def shift(h, deg): hh, s, l = hls(h); r, g, b = colorsys.hls_to_rgb(((hh + deg) % 360) / 360, l, s); return "#%02x%02x%02x" % (round(r * 255), round(g * 255), round(b * 255))
for pid in ["human", "vampire", "were", "ghost", "undead", "forest", "eldritch", "magic"]:
    p = PEO[pid]; c = p["color"]; hh, s, l = hls(c)
    ws.append([p["name"], c, p["colorName"], round(hh), round(s, 2), round(l, 2), shift(c, 180), shift(c, 150), shift(c, 210), shift(c, 30), shift(c, -30), MOOD[pid][0], MOOD[pid][1], MOOD[pid][2]])
style(ws, [14, 11, 12, 6, 9, 9, 11, 11, 11, 11, 11, 44, 44, 44], wrap=(11, 12, 13), swatch=(1, 6, 7, 8, 9, 10), freeze="C2")

ws = wb.create_sheet("MVP Trios")
ws.append(["#", "People", "Character", "People color", "Main (60%)", "Secondary (30%)", "Pop (10%)", "Main H", "Main S", "Main L", "Sec H", "Sec S", "Sec L", "Pop H", "Pop S", "Pop L", "Value spread", "Main vs Secondary contrast", "Pop vs Main contrast", "Min hue gap to people color", "Checks passed (of 5)", "Flags", "Ideas"])
for s in REAL:
    pc = s["peopleColor"]; m, sc, po = s["pal"]; f = []; ideas = []; passed = 0
    lm, ls_, lp = lum(m), lum(sc), lum(po); spread = max(lm, ls_, lp) - min(lm, ls_, lp)
    cms = contrast(m, sc); cpm = contrast(po, m); mh, ms, ml = hls(m); sh, ss, sl = hls(sc); ph, ps, pl = hls(po); gap = min(hdist(hls(pc)[0], x) for x in (mh, sh, ph) if True)
    # hue gap only counts for chromatic colors
    chroma = [(h, sat) for h, sat in ((mh, ms), (sh, ss), (ph, ps)) if sat > 0.2]
    gap = min([hdist(hls(pc)[0], h) for h, _ in chroma] or [180]) if hls(pc)[1] >= 0.2 else 180  # a near-neutral people color has no hue to clash with
    if spread >= 0.25: passed += 1
    else: f.append("value spread too small: the character may not read in grayscale"); ideas.append("Darken or lighten one accent")
    if cms >= 1.5 or hdist(mh, sh) >= 40: passed += 1
    else: f.append("main and secondary are too alike"); ideas.append("Move the secondary %s" % ("lighter" if lm < 0.4 else "darker") + " or shift its hue 40+ degrees")
    if ps >= 0.55 or cpm >= 2: passed += 1
    else: f.append("pop does not pop"); ideas.append("Raise pop saturation or move it to the complement of main: " + shift(m, 180))
    if gap >= 25: passed += 1
    else: f.append("an accent is within 25 degrees of the people color and may melt into the background"); ideas.append("Shift that accent's hue or pick a neutral for it")
    muddy = [x for x, (hs, ss2, ll) in zip(("main", "secondary", "pop"), ((mh, ms, ml), (sh, ss, sl), (ph, ps, pl))) if ss2 < 0.18 and 0.25 < ll < 0.75]
    if not muddy: passed += 1
    else: f.append("muddy mid-tone: " + ", ".join(muddy)); ideas.append("Either make it a clear neutral (very light or very dark) or add saturation")
    ws.append([s["n"], PEO[s["people"]]["name"], s["name"], pc, m, sc, po, round(mh), round(ms, 2), round(ml, 2), round(sh), round(ss, 2), round(sl, 2), round(ph), round(ps, 2), round(pl, 2), round(spread, 2), round(cms, 2), round(cpm, 2), round(gap), passed, "; ".join(f) or "none", "; ".join(ideas)])
style(ws, [4, 14, 30, 11, 11, 11, 11, 7, 7, 7, 7, 7, 7, 7, 7, 7, 9, 12, 12, 12, 10, 60, 70], wrap=(21, 22), swatch=(3, 4, 5, 6), freeze="D2")
n = ws.max_row
ws.conditional_formatting.add(f"U2:U{n}", CellIsRule(operator="greaterThanOrEqual", formula=["5"], fill=fill("C6EFCE")))
ws.conditional_formatting.add(f"U2:U{n}", CellIsRule(operator="between", formula=["3", "4"], fill=fill("FFEB9C")))
ws.conditional_formatting.add(f"U2:U{n}", CellIsRule(operator="lessThan", formula=["3"], fill=fill("FFC7CE")))

ws = wb.create_sheet("Skin Chart")
ws.append(["Row", "Light tone", "Shadow tone", "Use for"])
USE = ["very fair, cool-lit skin", "light to medium skin", "medium to tan skin", "brown skin", "deep brown skin", "very deep brown skin"]
for i, (a, b) in enumerate(A["skin"], 1): ws.append([i, a, b, USE[i - 1]])
style(ws, [6, 14, 14, 40], swatch=(1, 2), freeze="A2")

ws = wb.create_sheet("Shape and Material")
ws.append(["Element", "Rule", "Prompt wording (never reword)"])
for e in A["elements"]:
    if e["id"] in ("E1", "E2", "E3", "E4", "E5", "E6", "E7", "E8"): ws.append([f"{e['id']} {e['name']}", e["rule"], e["token"]])
ws.append(["", "", ""])
ws.append(["MATERIAL SHORTCUTS", "How to draw it", ""])
for r in next(e for e in A["elements"] if e["id"] == "E9")["table"]: ws.append([r[0], r[1], ""])
ws.append(["", "", ""]); ws.append(["PROPORTIONS BY SIZE", "", ""])
for k, v in A["sizes"].items(): ws.append([k.title(), v, ""])
ws.append(["", "", ""]); ws.append(["POSE ENERGY BY TYPE", "", ""])
for k, v in A["poses"].items(): ws.append([k.title(), v, ""])
style(ws, [26, 90, 110], wrap=(1, 2), freeze="A2")

ws = wb.create_sheet("Inspiration")
ws.append(["People", "Mood words", "Shape motifs", "Light and scene", "Color story", "Search cues (describe, don't name)"])
SH = {
    "human": ("wrench and gear wedges, riveted plates, antenna masts, boxy trapezoid armor", "rain, sodium lamps, shipping-container walls", "gold wash, teal and magenta accents, cream highlights", "street musicians in neon, mechanics and drifters, jumpsuits with patches, rooftop radio antenna silhouettes"),
    "vampire": ("tall collars, wing-cloak points, chandelier spikes, long thin faces", "ruined penthouse, one pale rim light, a huge moon", "ink black wash, violet and crimson accents, bone white", "gothic architecture silhouettes, bat swarms as shard clusters, dramatic high-contrast spot blacks"),
    "were": ("torn-edge fur shards, spiked collars, tire arcs, bone plates", "campfire barrels, scrap walls, painted totems", "rust wash, bone cream and hot pink pops", "heavy beast silhouettes, chunky claws, armor made of salvage"),
    "ghost": ("glitch blocks, cable curls, hollow masks, hanging coats", "dead CRT walls, one flickering screen, fog", "ice blue wash, hot magenta or orange pops for contrast", "broken signal patterns drawn as flat blocks, empty-eyed masks, trailing cloth"),
    "undead": ("ribbed bone arcs, stamps and seals, scythe blades, robe shards", "bunker corridors, filing walls, dusty shafts", "bone wash, toxic green and magenta pops, navy shadows", "skeleton silhouettes with bold trim, official-looking sigils, long hooded cloaks"),
    "forest": ("leaf wedges, vine spirals, antler forks, mossy plates", "overgrown streets, roots through concrete, mist", "moss green wash, red and gold accents", "creature-and-rider compositions, foliage as stacked shards, armored plant people"),
    "eldritch": ("curling tendrils, rows of teeth, insect wing panels, wrong angles", "flooded cathedral, black water, impossible stairs", "teal wash, crimson, violet and gold pops", "many-eyed masses, strange geometry, huge ragged wings"),
    "magic": ("pointed hats, staff heads, floating jars, spore clouds as shard clusters", "chalk circles, mushroom factory floors, a leaning observatory", "violet wash, lime, gold and pink accents", "silly-but-dangerous wizard poses, bubbling jars, conjuring hands thrust toward the viewer"),
}
for pid in ["human", "vampire", "were", "ghost", "undead", "forest", "eldritch", "magic"]:
    sh = SH[pid]; ws.append([PEO[pid]["name"], MOOD[pid][0], sh[0], sh[1], sh[2], sh[3]])
ws.append(["", "", "", "", "", ""]); ws.append(["BY ELEMENT", "What to look for when gathering references", "", "", "", ""])
for e, cue in [("E1 Shape language", "Characters whose silhouette is a cluster of spikes and wedges. Hair and capes ending in points."), ("E2 Proportions", "Tall tapered bodies with small heads and very big hands and boots."), ("E3 Perspective push", "Poses where a fist, foot or weapon comes straight at the camera and is huge."), ("E4 Ink line", "Black outlines that thicken on the underside and end in a flick."), ("E5 Flat fill", "Big areas of solid color with no blending at all."), ("E6 Two-tone shading", "One hard shadow shape per color and pure black in the creases."), ("E7 Face", "Big angular eyes, strong brows, wide grinning mouths with sharp teeth."), ("E8 Costume detail", "Heroic costumes reduced to three bold shapes: chunky belt, boots, shoulder plates.")]:
    ws.append([e, cue, "", "", "", ""])
style(ws, [18, 44, 60, 44, 44, 70], wrap=(1, 2, 3, 4, 5), freeze="A2")
wb.save(os.path.join(root, "NDA_Art_Color_Reference.xlsx"))
print("wrote NDA_MVP_Characters_V2.xlsx and NDA_Art_Color_Reference.xlsx")
